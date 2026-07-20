"""End-to-end structuring over a real generated .docx fixture (offline)."""
import os

os.environ.setdefault("OPENAI_API_KEY", "sk-test-key")

import asyncio
import io
import json

import pytest
from unittest.mock import AsyncMock, patch

import app.routes.documents as docs_mod
from app.conversion import LocalMarkItDownConverter
from app.structuring import structure_document

docx = pytest.importorskip("docx", reason="python-docx required for the DOCX fixture")

# python-docx has no high-level API for "repeat as header row" (w:tblHeader).
# Without it, plain python-docx tables have no row marked as a header, and
# mammoth (the DOCX->HTML step inside MarkItDown) then emits <td> for every
# row; MarkItDown's HTML->Markdown step still has to satisfy GFM's mandatory
# header-row syntax, so it invents a *blank* markdown header row and pushes
# what was visually the header ("leave_type"/"annual_days") down into the
# first data row -- losing the column names entirely. Setting w:tblHeader
# directly via oxml (a documented python-docx recipe for this exact gap) is
# what makes mammoth emit a real <th> row, which MarkItDown then renders as
# the markdown header `| leave_type | annual_days |`. Verified by probing
# the real conversion output before writing this fixture.
from docx.oxml.shared import OxmlElement


def _mark_header_row(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


def _make_docx_bytes() -> bytes:
    d = docx.Document()
    d.add_heading("Company Policy", level=1)
    d.add_heading("Leave", level=2)
    d.add_paragraph("Employees accrue leave each month.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "leave_type"
    table.rows[0].cells[1].text = "annual_days"
    _mark_header_row(table.rows[0])
    for lt, ad in [("Parental", "42"), ("Sick", "10")]:
        cells = table.add_row().cells
        cells[0].text = lt
        cells[1].text = ad
    d.add_heading("Pay", level=2)
    d.add_paragraph("Salaries are paid monthly.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_structuring_splits_headings_and_extracts_table():
    md = asyncio.run(
        LocalMarkItDownConverter(timeout_seconds=30).to_markdown(_make_docx_bytes(), "policy.docx")
    )
    doc = structure_document(
        markdown=md, source="policy.docx", file_type="docx",
        ingested_at="2026-07-20T00:00:00Z", chunk_size=500, chunk_overlap=50,
    )
    # 1. narrative split on the heading hierarchy
    paths = {tuple(c.heading_path) for c in doc.narrative_chunks}
    assert any("Leave" in p for p in paths)
    assert any("Pay" in p for p in paths)
    # 2. table rows extracted separately (not inline in narrative)
    assert len(doc.table_chunks) >= 1
    rows = doc.table_chunks[0].rows
    assert {"leave_type": "Parental", "annual_days": "42"} in rows
    assert all("| Parental |" not in c.text for c in doc.narrative_chunks)


@pytest.mark.asyncio
async def test_docx_upload_updates_manifest(client, monkeypatch, tmp_path):
    monkeypatch.setattr(docs_mod.settings, "converted_output_dir", str(tmp_path))
    with (
        patch.object(docs_mod, "embed_texts", new_callable=AsyncMock,
                     side_effect=lambda texts: [[0.0, 0.0, 0.0] for _ in texts]),
        patch.object(docs_mod, "upsert_chunks", new_callable=AsyncMock),
    ):
        files = {"file": ("policy.docx", _make_docx_bytes(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = await client.post("/documents/upload", files=files)
    assert r.status_code == 201

    data = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    assert "policy.docx" in data
    assert len(data["policy.docx"]["chunk_ids"]) >= 1
    assert len(data["policy.docx"]["table_ids"]) >= 1
